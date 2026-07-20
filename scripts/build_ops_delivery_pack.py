from __future__ import annotations

import argparse
import json
import sys
from datetime import datetime
from pathlib import Path
from typing import Any

from local_demand_radar import scan_paths, write_outputs


BRAND_BG = "050816"
BRAND_SURFACE = "0F172A"
BRAND_TEXT = "F8FAFC"
BRAND_MUTED = "94A3B8"
BRAND_PINK = "FF3D9A"
BRAND_CYAN = "22D3EE"
BRAND_GOLD = "FACC15"


def ensure_utf8() -> None:
    try:
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    except Exception:
        pass


def write_text(path: Path, content: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(content.strip() + "\n", encoding="utf-8")


def build_service_script_library(out_dir: Path) -> Path:
    content = """
# AI客服脚本型应答库

## 使用原则

- 网页客服：可以自动回复，所有会话落库。
- 微信/抖音/淘宝/拼多多/闲鱼：没有官方 API 前，只做读取、生成候选回复、人工确认发送。
- 退款、投诉、账号、付款、隐私、辱骂：自动标记人工跟进，回复要谨慎。
- 不承诺绝对准时、绝对收益、百分百转化、无条件退款。

## 通用新客开场

客户：这个怎么做？

回复：
您好，我先帮您确认一下需求。您现在主要是想解决「客服没人回、获客断了、还是内容素材做不出来」哪一块？我可以按场景给您先配一套能测试的方案。

下一步：
请客户发行业、当前平台、每天咨询量、最想自动化的动作。

## AI客服成交脚本

客户：你这个 AI 客服能不能全自动？

回复：
网页客服可以直接全自动回复，并且客户消息、AI回复、线索状态都会保存到后台。微信、抖音、淘宝、拼多多、闲鱼这些平台，如果要自动发送，需要官方 API 或开放平台权限；没有权限前我们先做“自动读取 + 自动生成回复 + 人工确认发送”，这样稳定也更安全。

## 抖音小店/私信说明

客户：能不能打开抖音小店就自动回？

回复：
可以做成抖音小店客服接入方案，但要分两步：第一步先把商品、售后、价格、常见问题导入知识库，系统生成回复草稿；第二步拿到抖音小店/飞鸽/开放平台相关权限后，再接官方接口做自动回复。我们不做绕过风控的模拟发送。

## 价格异议

客户：太贵了。

回复：
理解，前期最重要不是一下子买大系统，而是先把一个能赚钱的场景跑通。可以先从“网页客服 + 客服话术库 + 线索表”开始，小范围测试有咨询、有回复、有留资，再决定是否扩展到抖音/微信/电商平台。

## 交付边界

客户：今天能全部做好吗？

回复：
今天可以先交付可验收版本：后台可登录、网页客服可自动回复、知识库可导入、桌面助手可生成回复草稿、客户线索表和运营方案可使用。平台官方自动发送需要账号资质和 API 权限，权限到位后再接。

## 高风险兜底

客户：投诉/退款/付款/账号/隐私相关。

回复：
这类问题我先帮您记录并转人工处理，避免信息误差。麻烦您留下订单号/手机号/具体问题，我们会尽快核实后回复。
"""
    path = out_dir / "01_AI客服脚本型应答库.md"
    write_text(path, content)
    return path


def build_operator_guide(out_dir: Path) -> Path:
    content = """
# 今日可商用操作指南

## 先用哪种方式

优先用网页端客服。原因：

- 最稳定：不依赖第三方平台窗口是否可读。
- 风控最低：客户主动在你的网页咨询，AI 自动回复属于自有系统能力。
- 可落库：客户消息、AI 回复、线索状态都能在后台查看。

抖音开发者平台、小程序、抖音小店官方接口放第二阶段。它们更适合上架和长期运营，但需要资质、审核、权限和接口调试。

## 一步一步使用

1. 打开后台：https://wjhai.cn/merchant-admin/
2. 登录：ai_kefu_demo / admin123
3. 进入“知识库导入”，粘贴商品、价格、FAQ、售后规则。
4. 进入“客服脚本”，生成新客、价格异议、售后、留资脚本。
5. 进入“网页气泡”，复制 script 到网站或打开测试页。
6. 进入“会话收件箱”，查看客户消息、AI 回复和待跟进线索。

## 桌面辅助客服怎么测

1. 打开微信/抖音/千牛/拼多多/闲鱼真实客户聊天页。
2. 双击：启动AI自动客服.bat
3. 选择目标窗口。
4. 默认只粘贴候选回复，不自动按 Enter。
5. 运行“验收真实平台.bat”检查窗口是否真的可读。

## 抖音小店自动回复怎么落地

当前合规路线：

- 没有官方权限：只做桌面辅助回复草稿。
- 有官方权限：接抖音小店/飞鸽/开放平台消息接口，AI 回复后由官方接口发送。
- 不做：绕过权限的模拟点击、批量骚扰私信、抓取联系方式。

## 今日验收

- npm run acceptance:check
- npm run acceptance:reply-e2e
- npm run demand:scan
- npm run ops:package
- 打开真实聊天页后运行：npm run acceptance:real-platforms
"""
    path = out_dir / "02_今日可商用操作指南.md"
    write_text(path, content)
    return path


def build_compliance_sop(out_dir: Path) -> Path:
    content = """
# 合规获客与全平台推流 SOP

## 今日先做

- 抖音企业号主页挂服务介绍和线索入口。
- 每天 3 条内容：AI客服演示、商家断客问题、网页客服自动回复案例。
- 评论区只引导主动咨询或表单留资，不批量私信骚扰。
- 将主动咨询、表单、微信添加、闲鱼留言统一录入客户线索表。

## 不能做

- 不能无授权抓取手机号、微信号、联系方式。
- 不能批量自动私信陌生人。
- 不能绕过平台风控模拟发送。

## 可以做

- 手动搜索潜在商家，记录公开主页、行业、需求判断。
- 使用平台广告/企业号线索表单收集客户主动提交的信息。
- 对已咨询客户做自动分级、话术生成和跟进提醒。

## 公司注册后矩阵

- 抖音：企业号 + 线索表单 + 同城案例 + 小店客服方案。
- 微信：公众号/企微/社群承接，客服系统作为官网入口。
- 闲鱼：卖“AI客服部署/小程序/运营系统”服务，承接主动咨询。
- 小红书：发案例笔记和模板前后对比。
- B站/视频号：发长教程建立信任。
- GitHub：开源核心脚本，卖私有部署、二开、模板库。
"""
    path = out_dir / "03_合规获客与推流SOP.md"
    write_text(path, content)
    return path


def build_api_checklist(out_dir: Path) -> Path:
    content = """
# 官方 API 接入清单

## 已能直接使用

- 网页客服气泡：自动回复、会话落库、线索标记。
- API2D/兼容 OpenAI 中转：用于生成回复。
- 桌面辅助：读取真实聊天页，生成候选回复，人工确认发送。

## 抖音小店/飞鸽

需要：

- 抖音开放平台或小店服务市场相关资质。
- 消息事件回调权限。
- 发送客服消息权限。
- 店铺授权流程和回调地址。

未拿到权限前：

- 不承诺全自动发送。
- 只做话术库、回复草稿、人工确认。

## 微信

可选路线：

- 微信客服官方接口。
- 企业微信客服。
- 公众号客服消息。

个人微信不建议做全自动模拟发送，容易触发限制。

## 淘宝/拼多多/闲鱼

优先路线：

- 千牛/淘宝开放平台。
- 拼多多开放平台/商家客服能力。
- 闲鱼属于淘宝生态，先做桌面辅助，官方能力到位后接入。
"""
    path = out_dir / "04_官方API接入清单.md"
    write_text(path, content)
    return path


def build_workbook(out_dir: Path, demand_rows: list[Any]) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "今日执行看板"
    ws.append(["模块", "今日状态", "验收方式", "负责人/备注"])
    rows = [
        ["AI客服后台", "可用", "登录后台 + acceptance:check", "优先用网页端"],
        ["网页客服气泡", "可用", "打开 widget-test 测消息", "可自动回复"],
        ["桌面辅助客服", "需真实聊天页", "acceptance:real-platforms", "不做违规自动发送"],
        ["抖音小店自动回复", "待官方 API 权限", "填写开放平台权限清单", "先做草稿回复"],
        ["客户线索表", "可用", "手动/授权导入", "禁止无授权抓取联系方式"],
        ["需求雷达", "可用", "npm run demand:scan", "扫描本项目文件"],
        ["PPT/方案", "可用", "查看交付包", "用于谈客户"],
    ]
    for row in rows:
        ws.append(row)

    lead = wb.create_sheet("客户线索表")
    lead.append(["录入日期", "平台", "客户主页/来源链接", "客户名称", "行业", "公开账号ID", "联系方式(客户主动提供)", "需求", "意向分", "跟进状态", "下一步", "备注"])
    sample_leads = [
        [datetime.now().strftime("%Y-%m-%d"), "抖音", "客户主动咨询/线索表单", "示例商家A", "本地生活", "", "", "想解决无人回复和断客", 85, "待沟通", "发网页客服演示链接", "只记录主动咨询信息"],
        [datetime.now().strftime("%Y-%m-%d"), "闲鱼", "商品留言", "示例商家B", "电商", "", "", "想做AI客服和主图视频", 78, "待报价", "发Starter/Pro方案", "不要批量私信陌生人"],
    ]
    for row in sample_leads:
        lead.append(row)

    demand = wb.create_sheet("需求雷达Top")
    demand.append(["分数", "分类", "标题", "文件", "证据", "建议", "变现方式"])
    for item in demand_rows[:50]:
        demand.append([item.score, item.category, item.title, item.path, item.evidence, item.recommendation, item.monetization])

    channels = wb.create_sheet("平台接入状态")
    channels.append(["平台", "今日可用方式", "全自动发送条件", "风险边界"])
    channel_rows = [
        ["网页客服", "script 气泡自动回复", "已具备", "自有网站，风险最低"],
        ["微信", "桌面辅助草稿", "微信客服/企微官方接口", "个人微信不做模拟发送"],
        ["抖音", "桌面辅助草稿/企业号表单", "开放平台/小店客服权限", "不批量私信陌生人"],
        ["淘宝/千牛", "桌面辅助草稿", "淘宝开放平台/千牛能力", "不绕过平台控件"],
        ["拼多多", "桌面辅助草稿", "拼多多开放平台", "需商家端权限"],
        ["闲鱼", "桌面辅助草稿", "淘宝生态官方能力", "先承接主动咨询"],
    ]
    for row in channel_rows:
        channels.append(row)

    for sheet in wb.worksheets:
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor=BRAND_SURFACE)
            cell.font = Font(color=BRAND_TEXT, bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for column in range(1, sheet.max_column + 1):
            width = 16
            if column in {3, 4, 5, 6, 7}:
                width = 28
            sheet.column_dimensions[get_column_letter(column)].width = width
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    path = out_dir / "客户线索与需求运营表.xlsx"
    wb.save(path)
    return path


def add_textbox(slide: Any, left: float, top: float, width: float, height: float, text: str, size: int = 20, color: str = BRAND_TEXT, bold: bool = False) -> Any:
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor

    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.clear()
    p = frame.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    return box


def build_ppt(out_dir: Path) -> Path:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slides = [
        ("AI全自动化运营落地方案", "先跑通客服承接，再扩展获客、内容、小程序和官方 API。"),
        ("今日优先级", "1. 网页客服自动回复\n2. 客服脚本库\n3. 合规线索表\n4. 文件需求雷达\n5. 小程序/视频生成后续模块"),
        ("为什么先用网页客服", "稳定、低风控、可落库、可演示、可复制给商家。抖音开发者平台适合作为第二阶段上架入口。"),
        ("AI客服闭环", "客户进入网页/小程序/私域入口 -> AI自动回复 -> 高意向标记 -> 人工跟进 -> 客户线索表 -> 成交复盘。"),
        ("抖音小店路线", "无官方权限：桌面辅助草稿。\n有官方权限：接消息回调和客服发送接口。\n不做绕过风控的模拟发送。"),
        ("全网获客矩阵", "抖音企业号、小红书案例、闲鱼服务商品、微信私域、GitHub开源、官网SEO、视频号教程。"),
        ("可卖的AI项目", "AI客服SaaS、客服知识库部署、主图详情图生成、短视频画布、线索表/CRM、小程序模板。"),
        ("今天验收标准", "后台可登录、网页客服能回、知识库能导入、脚本能生成、线索表/PPT/需求雷达可交付、真实聊天页可诊断。"),
    ]

    for i, (title, body) in enumerate(slides):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        bg = slide.background.fill
        bg.solid()
        bg.fore_color.rgb = RGBColor.from_string(BRAND_BG)
        accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.16))
        accent.fill.solid()
        accent.fill.fore_color.rgb = RGBColor.from_string(BRAND_PINK if i % 2 == 0 else BRAND_CYAN)
        accent.line.fill.background()
        add_textbox(slide, 0.7, 0.62, 11.8, 0.8, title, 38, BRAND_TEXT, True)
        add_textbox(slide, 0.72, 1.55, 10.8, 3.6, body, 25, BRAND_TEXT, False)
        chip = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.72), Inches(6.35), Inches(3.7), Inches(0.46))
        chip.fill.solid()
        chip.fill.fore_color.rgb = RGBColor.from_string(BRAND_SURFACE)
        chip.line.color.rgb = RGBColor.from_string(BRAND_CYAN)
        add_textbox(slide, 0.88, 6.42, 3.2, 0.3, "Merchant AI Ops System", 13, BRAND_CYAN, True)

    path = out_dir / "AI全自动化运营落地方案.pptx"
    prs.save(path)
    return path


def build_pdf(out_dir: Path) -> Path:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    font_name = "Helvetica"
    for candidate in [Path(r"C:\Windows\Fonts\msyh.ttc"), Path(r"C:\Windows\Fonts\simsun.ttc"), Path(r"C:\Windows\Fonts\simhei.ttf")]:
        if candidate.exists():
            try:
                pdfmetrics.registerFont(TTFont("CNFont", str(candidate)))
                font_name = "CNFont"
                break
            except Exception:
                continue

    path = out_dir / "AI客服商用操作手册.pdf"
    doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm, bottomMargin=16 * mm)
    styles = getSampleStyleSheet()
    title = ParagraphStyle("TitleCN", parent=styles["Title"], fontName=font_name, fontSize=22, leading=28, textColor="#111827")
    h2 = ParagraphStyle("H2CN", parent=styles["Heading2"], fontName=font_name, fontSize=14, leading=20, textColor="#111827")
    body = ParagraphStyle("BodyCN", parent=styles["BodyText"], fontName=font_name, fontSize=10.5, leading=16, textColor="#111827")

    story = [
        Paragraph("AI客服商用操作手册", title),
        Spacer(1, 8),
        Paragraph("1. 先用网页客服。复制后台接入代码到网站，客户咨询后自动回复并保存会话。", body),
        Paragraph("2. 微信、抖音、淘宝、拼多多、闲鱼先用桌面辅助草稿。没有官方权限前不做自动发送。", body),
        Paragraph("3. 每天把主动咨询、表单线索、闲鱼留言录入客户线索表，按意向分跟进。", body),
        Spacer(1, 12),
        Paragraph("今日验收命令", h2),
        Paragraph("npm run acceptance:check<br/>npm run acceptance:reply-e2e<br/>npm run demand:scan<br/>npm run ops:package", body),
        Spacer(1, 12),
        Paragraph("抖音小店接入边界", h2),
        Paragraph("抖音小店要做到全自动回复，需要官方消息回调和客服发送权限。权限未到位时，只能做桌面辅助和回复草稿。", body),
        Spacer(1, 12),
        Paragraph("客户成交路径", h2),
        Paragraph("内容引流 -> 表单/私信主动咨询 -> AI客服承接 -> 高意向标记 -> 人工跟进 -> 试用/部署成交。", body),
    ]
    doc.build(story)
    return path


def main() -> int:
    ensure_utf8()
    parser = argparse.ArgumentParser(description="Build today's AI operations delivery pack.")
    parser.add_argument("--output-dir", default="运营计划/今日交付包", help="Output directory.")
    args = parser.parse_args()

    root = Path(__file__).resolve().parents[1]
    out_dir = root / args.output_dir
    out_dir.mkdir(parents=True, exist_ok=True)

    hits = scan_paths(root, ["docs", "运营计划", "product-kit", "douyin-miniapp", "src", "backend", "scripts"])
    demand_outputs = write_outputs(root, hits[:200], out_dir / "需求雷达")
    files = {
        "script_library": str(build_service_script_library(out_dir)),
        "operator_guide": str(build_operator_guide(out_dir)),
        "compliance_sop": str(build_compliance_sop(out_dir)),
        "api_checklist": str(build_api_checklist(out_dir)),
        "workbook": str(build_workbook(out_dir, hits)),
        "ppt": str(build_ppt(out_dir)),
        "pdf": str(build_pdf(out_dir)),
        "demand": demand_outputs,
    }
    readme = [
        "# 今日交付包",
        "",
        "先看顺序：",
        "",
        "1. `02_今日可商用操作指南.md`",
        "2. `01_AI客服脚本型应答库.md`",
        "3. `客户线索与需求运营表.xlsx`",
        "4. `AI全自动化运营落地方案.pptx`",
        "5. `需求雷达/需求雷达报告.md`",
        "",
        "重要边界：抖音小店、微信、淘宝、拼多多、闲鱼要全自动发送，必须接官方 API 权限；今天可商用的是网页客服自动回复和桌面辅助草稿。",
    ]
    readme_path = out_dir / "先看这个-今日交付包.md"
    write_text(readme_path, "\n".join(readme))
    files["readme"] = str(readme_path)
    print(json.dumps({"ok": True, "output_dir": str(out_dir), "files": files}, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
