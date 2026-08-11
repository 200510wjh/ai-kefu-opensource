from pathlib import Path
from docx import Document
from pptx import Presentation
from openpyxl import load_workbook

FILES = [
    r"D:\xwechat_files\wxid_x74btth050jk22_cf27\msg\file\2026-04\Recardify智能客服配置方案.md",
    r"D:\xwechat_files\wxid_x74btth050jk22_cf27\msg\file\2026-04\副本-王鲜记产品知识+客服话术.xlsx",
    r"D:\xwechat_files\wxid_x74btth050jk22_cf27\msg\file\2026-05\小危AI电信行业智能客服方案-修改版-2.pptx",
    r"D:\xwechat_files\wxid_x74btth050jk22_cf27\msg\file\2026-05\陈-timeless-电信行业AI客服方案提纲(1)(1).docx",
    r"D:\xwechat_files\wxid_x74btth050jk22_cf27\msg\file\2026-05\陈-timeless-电信行业AI客服方案提纲(2).docx",
    r"D:\xwechat_files\wxid_x74btth050jk22_cf27\msg\file\2026-06\股东合作协议草案_AI客服公司.docx",
    r"D:\xwechat_files\wxid_x74btth050jk22_cf27\msg\file\2026-06\股东合作协议草案_AI客服公司(1).docx",
    r"D:\xwechat_files\wxid_x74btth050jk22_cf27\msg\file\2026-07\陈-timeless-电信行业AI客服方案提纲.docx",
]

OUT = Path(r"C:\Users\Administrator\Documents\运营\work\ai_customer_replay\wechat_file_extract.txt")


def extract_docx(path):
    doc = Document(path)
    parts = []
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt:
            parts.append(txt)
    for table in doc.tables:
        for row in table.rows:
            vals = [c.text.strip().replace("\n", " / ") for c in row.cells]
            if any(vals):
                parts.append(" | ".join(vals))
    return "\n".join(parts)


def extract_pptx(path):
    prs = Presentation(path)
    parts = []
    for i, slide in enumerate(prs.slides, 1):
        slide_text = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                slide_text.append(shape.text.strip().replace("\n", " / "))
        if slide_text:
            parts.append(f"[Slide {i}] " + " || ".join(slide_text))
    return "\n".join(parts)


def extract_xlsx(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append(f"[Sheet] {ws.title}")
        for row in ws.iter_rows(max_row=40, values_only=True):
            vals = [str(v).strip() for v in row if v is not None and str(v).strip()]
            if vals:
                parts.append(" | ".join(vals))
    return "\n".join(parts)


def extract_text(path):
    p = Path(path)
    if not p.exists():
        return "[missing]"
    suffix = p.suffix.lower()
    if suffix == ".docx":
        return extract_docx(p)
    if suffix == ".pptx":
        return extract_pptx(p)
    if suffix in [".xlsx", ".xlsm"]:
        return extract_xlsx(p)
    if suffix in [".md", ".txt"]:
        return p.read_text(encoding="utf-8", errors="ignore")
    return "[unsupported]"


def main():
    chunks = []
    for f in FILES:
        chunks.append("=" * 90)
        chunks.append(f)
        chunks.append("-" * 90)
        text = extract_text(f)
        chunks.append(text[:12000])
    OUT.write_text("\n".join(chunks), encoding="utf-8")
    print(OUT)


if __name__ == "__main__":
    main()
