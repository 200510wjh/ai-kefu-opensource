from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from build_full_scan_report import completion_audit_rows, contacts, customer_tag_rows, file_evidence_rows, products


ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "outputs" / "AI客服客户扫描台账.xlsx"


def style_sheet(ws, widths):
    header_fill = PatternFill("solid", fgColor="1F4D78")
    header_font = Font(name="Microsoft YaHei", color="FFFFFF", bold=True)
    body_font = Font(name="Microsoft YaHei", size=10)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = body_font
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = width


def add_sheet(wb, title, headers, rows, widths):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for row in rows:
        ws.append(row)
    style_sheet(ws, widths)
    return ws


def classify(row):
    name, source, evidence, scene, judgement, next_step = row
    if evidence == "聊天证据":
        level = "A-聊天已验证"
    elif "文件" in evidence:
        level = "A-文件已验证"
    elif "客户标签" in evidence or "通讯录" in evidence:
        level = "B-标签/备注证据"
    else:
        level = "C-待补"

    text = f"{source} {scene} {judgement}"
    if any(k in text for k in ["王鲜记", "王超", "三w", "岁月留痕", "抖店", "淘宝", "生鲜", "商品"]):
        product = "多渠道电商客服包 / 商品知识库+话术包"
    elif any(k in text for k in ["企业微", "企微", "企业微信"]):
        product = "低风险企微辅助包"
    elif any(k in text for k in ["saas", "SaaS", "系统"]):
        product = "SaaS系统嵌入包"
    elif any(k in text for k in ["电信", "VOS", "WhatsApp", "工单"]):
        product = "电信企业级集成包"
    elif any(k in text for k in ["关键词", "Daisy"]):
        product = "关键词客服升级包"
    elif any(k in text for k in ["图片", "内容", "超级员工", "运营"]):
        product = "AI内容/运营扩展包"
    else:
        product = "待确认后匹配"

    if name in {"王超", "王鲜记", "三w", "岁月留痕", "大头哥", "陈-timeless"}:
        priority = "P0/P1"
    elif "待判定" in judgement or "未验证" in judgement:
        priority = "P3"
    else:
        priority = "P2"
    return [name, level, priority, product, source, scene, judgement, next_step]


def main():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)

    add_sheet(
        wb,
        "客户标签19人",
        ["序号", "联系人", "备注/线索", "标签"],
        customer_tag_rows,
        [8, 18, 42, 18],
    )

    ledger_rows = [classify(row) for row in contacts]
    add_sheet(
        wb,
        "AI客户扫描台账",
        ["客户/资产", "证据级别", "优先级", "推荐产品包", "标签/来源", "需求或场景", "判断", "下一步"],
        ledger_rows,
        [22, 18, 10, 26, 28, 36, 40, 42],
    )

    add_sheet(
        wb,
        "微信文件证据",
        ["文件/资产", "看到的内容", "关键交付件", "对下次成交的意义"],
        file_evidence_rows,
        [32, 42, 46, 46],
    )

    add_sheet(
        wb,
        "可复制产品包",
        ["产品包", "适合客户", "核心交付", "为什么先卖"],
        products,
        [30, 38, 46, 34],
    )

    add_sheet(
        wb,
        "跟进优先级",
        ["优先级", "对象", "为什么", "第一动作"],
        [
            ["P0", "王鲜记 / 王超", "已有系统、话术、报价或群内客服助理场景，是最容易做样板的资产", "整理成商品知识库+客服话术+插件包案例"],
            ["P0", "三w / 岁月留痕", "抖店/抖音客服场景与王超相似", "用王超案例问是否要先跑一个商品问答场景"],
            ["P1", "大头哥", "明确是SaaS系统AI客服，阻力是价格和部署", "先卖轻量嵌入启动包"],
            ["P1", "陈-timeless", "电信企业级高客单，方案资料完整", "发一页纸试点版，拆第一阶段"],
            ["P2", "Daisy / Bryce / Mr.C / 我是超人", "已有关键词系统或担心企微接口，需降低风险和证明增量", "分别卖关键词升级包或低风险企微辅助包"],
            ["P3", "木木 / 大魔王 / 魁星Ai备... / 我来依旧 / 张口就来 / Li.", "目前主要是标签或备注证据", "补聊天原文：行业、咨询产品、报价、最后一句"],
        ],
        [10, 34, 54, 48],
    )

    add_sheet(
        wb,
        "完成度审计",
        ["要求", "状态", "当前证据", "剩余缺口"],
        completion_audit_rows,
        [22, 16, 58, 42],
    )

    wb.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
