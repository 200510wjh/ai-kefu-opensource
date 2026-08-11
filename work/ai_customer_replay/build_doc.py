from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = r"C:\Users\Administrator\Documents\运营\outputs\AI客服客户洞察与成交复制手册.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(9.5)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_text(cell, header, bold=True, color="FFFFFF")
        set_cell_shading(cell, "1F4D78")
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], value)
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Inches(width)
    doc.add_paragraph()
    return table


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.name = "Arial"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.color.rgb = RGBColor.from_string("1F4D78" if level == 1 else "0B2545")
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(item)
        run.font.name = "Arial"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.size = Pt(10.5)


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Normal"].font.size = Pt(10.5)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("AI客服客户洞察与成交复制手册")
    r.bold = True
    r.font.size = Pt(22)
    r.font.name = "Arial"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.color.rgb = RGBColor.from_string("0B2545")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = subtitle.add_run("用于沉淀微信客户标签、复盘成交原因、提炼卖点，并复制下一批可成交产品")
    sr.font.size = Pt(10.5)
    sr.font.name = "Arial"
    sr._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    sr.font.color.rgb = RGBColor.from_string("555555")

    add_heading(doc, "1. 使用说明", 1)
    add_bullets(doc, [
        "资料范围：微信中客户标签包含 ai、AI、ai客服、AI客服、智能客服、数字人客服等关键词的联系人。",
        "读取维度：客户画像、业务场景、痛点、预算/决策、已购产品、未成交原因、下一步成交机会。",
        "输出原则：只把微信中能证实的信息写成结论；不确定内容标记为“待验证”，避免凭感觉做产品。",
    ])

    add_heading(doc, "2. 已从微信查看到的客户线索", 1)
    add_table(
        doc,
        ["客户/备注", "标签/场景", "微信里看到的原话或证据", "判断", "下一步动作"],
        [
            ["Bryce", "ai客服", "对方说：企微关闭这个接口的话暂时估计不上，等企微上线后再考虑；他们已有一套客服团队。", "未成交。核心阻力是接口风险和已有客服团队，AI不是刚需。", "不要继续讲功能，改讲低风险方案：不依赖企微接口的网页客服/辅助回复/人工兜底。"],
            ["Daisy", "ai客服客户", "对方说：已经在使用普通客服系统，通过关键词能解决问题，所以没有想弄AI。", "未成交。已有替代方案，AI增量价值不够清晰。", "用“关键词客服 vs AI客服”的对比演示，证明AI能处理模糊问法、追问和线索分层。"],
            ["陈-timeless", "ai客服", "对方说晚一点、这几天在出差；你后续追问“哥您那边怎么说”。", "跟进中。阻力不是产品，而是时间和响应断点。", "发一页纸总结和两个可选时间，让对方低成本继续。"],
            ["王超群", "抖音/淘宝/微信公众号ai客服", "群内客户问“六月黄有吗”“有螃蟹吗？@客服助理”，客服助理自动发名片并提示点击发起咨询。", "已验证/疑似已交付场景。多渠道电商客服助理可以复制。", "把它包装成“电商客服助理样板案例”，沉淀截图、流程、报价和复购包。"],
            ["陈-timeless", "电信行业AI客服", "微信文件有《陈-timeless-电信行业AI客服方案提纲.docx》，方案写明客户为陈-timeless、行业为电信服务。", "高价值售前客户。需求不是普通FAQ，而是企业级定制。", "单独作为“电信行业大单案例”推进，重点跟进决策人、预算、试点范围。"],
            ["小危AI电信方案", "企业级AI客服PPT", "微信文件有《小危AI电信行业智能客服方案-修改版-2.pptx》，包含痛点、六大模块、报价和ROI。", "已经形成成熟售前方案。", "把这套PPT复制成行业模板，适配运营商、呼叫中心、海外客服团队。"],
            ["王鲜记", "商品知识+客服话术/客服系统", "微信文件有《副本-王鲜记产品知识+客服话术.xlsx》《王鲜记客服系统.zip》《王鲜记AI内容生产服务报价单.md.pdf》。", "这是最接近“已交付产品包”的客户资产。", "复制为“电商商品知识库+客服话术+内容生产”组合包。"],
            ["Recardify", "智能客服配置方案", "微信文件有《Recardify智能客服配置方案.md》。", "有配置方案资产，说明客户需求偏“部署和配置”。", "复制为“海外/独立站智能客服配置包”，重点卖配置、话术、FAQ和渠道接入。"],
            ["佰社区", "AI超级员工", "微信文件有《佰社区AI超级员工.pdf》。", "需求可能不只是客服，而是AI员工/社区运营。", "拆成“AI客服+AI运营助理+线索沉淀”组合，不要只按客服卖。"],
            ["曹智腾", "AI客服方案", "搜索记录出现“AI客服方案提纲.docx”。", "已有方案交付痕迹，可能是售前方案或交付方案。", "继续定位原聊天上下文，补充客户行业、报价、成交状态。"],
            ["jack / Torre / 小危", "ai客服相关聊天记录", "搜索聊天记录中出现相关结果，但窗口最小化后未继续展开。", "待补充。", "恢复微信窗口后继续读取，补齐行业、需求和结果。"],
        ],
        [0.9, 1.05, 2.65, 1.55, 1.85],
    )

    add_heading(doc, "3. 从真实客户里抽出的需求和阻力", 1)
    add_table(
        doc,
        ["类型", "客户证据", "说明", "销售含义"],
        [
            ["接口风险", "Bryce 担心企微接口关闭", "客户不是不想自动化，而是怕买完不能稳定用。", "卖点要加“合规接入、替代通道、人工兜底、接口变更预案”。"],
            ["已有系统替代", "Daisy 已用关键词客服", "普通关键词能覆盖简单问题，AI必须证明新增收益。", "演示模糊问题、多轮追问、意向识别、线索沉淀。"],
            ["时间断点", "陈-timeless 出差后未继续", "意向客户容易因为忙而断掉，不一定是拒绝。", "用轻量资料和预约式跟进，而不是连续追问。"],
            ["已验证场景", "王超群里客服助理自动承接咨询", "电商生鲜/本地商品问答有明确可复制场景。", "优先复制到抖音、淘宝、公众号、微信群承接。"],
            ["企业级复杂需求", "陈-timeless 电信方案包含VOS、WhatsApp、App、工单、TTS/ASR、合规等模块", "大客户不是买一个聊天机器人，而是买系统集成和运营效率。", "报价不能按低价SaaS卖，要分阶段搭建费+月费+超量对话收费。"],
            ["知识库交付", "王鲜记有产品知识+客服话术表", "客户愿意为“把商品知识变成可回复系统”付费。", "标准化交付：商品资料表、FAQ、客服话术、自动回复、复盘报告。"],
        ],
        [0.95, 1.7, 2.0, 2.35],
    )

    add_heading(doc, "4. 已成交/可复制客户的共同规律", 1)
    add_table(
        doc,
        ["判断维度", "需要寻找的信号", "对产品复制的意义"],
        [
            ["痛点强度", "反复提到客服忙、漏单、回复慢、夜间无人、员工不稳定", "卖点应从“AI很先进”改成“少漏单、少人力、快响应”"],
            ["成交触发", "被案例、演示、低门槛试用、明确报价打动", "下次销售先给可视化演示，再谈套餐"],
            ["预算信号", "主动问价格、部署周期、是否包售后、能否分阶段做", "产品包要有入门版、标准版、增长版"],
            ["决策角色", "老板本人、运营负责人、客服主管、门店负责人", "不同角色分别强调利润、人效、稳定、可控"],
        ],
        [1.15, 2.55, 3.0],
    )

    add_heading(doc, "5. 未成交因素排查", 1)
    add_table(
        doc,
        ["阻力类型", "常见表现", "应对动作"],
        [
            ["需求不急", "先了解、以后再说、现在团队还能撑", "用漏单/响应慢的损失测算制造紧迫感"],
            ["信任不足", "担心 AI 不准、怕客户体验差", "给同行案例、试用账号、可人工兜底流程"],
            ["价格犹豫", "觉得贵、预算还没批", "拆成轻量启动包，先卖一个可验证场景"],
            ["交付不清", "不知道上线后谁维护、怎么训练知识库", "把交付步骤、验收标准、售后边界写清楚"],
            ["决策链断点", "对接人认可但老板没拍板", "准备老板版一页纸：收益、成本、风险、周期"],
        ],
        [1.05, 2.2, 3.45],
    )

    add_heading(doc, "6. 可复制产品包", 1)
    add_table(
        doc,
        ["产品包", "适合客户", "核心交付", "成交话术"],
        [
            ["客服助理启动包", "像 Daisy 一样已有关键词客服、但想提升承接质量的客户", "FAQ知识库、模糊问法识别、多轮追问、人工转接、7天调优", "不是替代你现有客服，是把关键词接不住的问题先补上。"],
            ["多渠道电商客服包", "像王超群一样有抖音、淘宝、公众号、微信群咨询的商家", "商品问答、名片/链接承接、咨询分流、客服助理提示、线索记录", "客户问有没有货、多少钱、怎么买，AI先接住，再交给人成交。"],
            ["低风险企微替代包", "像 Bryce 一样担心企微接口风险的客户", "网页客服/桌面辅助回复/人工确认发送/接口变更预案", "先不上高风险接口，先做能稳定跑的辅助客服。"],
            ["电信行业企业包", "像陈-timeless一样有呼叫中心、工单、App、WhatsApp需求的企业", "VOS电话对接、WhatsApp、App SDK、工单派发、质检验收、知识库", "先做核心AI客服和一个渠道试点，再扩展到电话、App和工单。"],
            ["商品知识库交付包", "像王鲜记一样商品多、问答重复、需要内容和客服一起交付的商家", "商品知识表、客服话术、FAQ、自动回复、内容生产报价", "先把商品知识整理成能成交的话术，再接自动客服。"],
            ["月度运营复盘包", "已经试用或上线AI客服但不知道效果的客户", "咨询统计、问题归类、话术优化、知识库更新、增购建议", "帮你看清 AI 少漏了多少咨询，哪些问题还该交给人。"],
        ],
        [1.2, 1.8, 2.0, 2.0],
    )

    add_heading(doc, "7. 下一次直接销售流程", 1)
    add_bullets(doc, [
        "第一步：用客户原话确认痛点，不先讲功能。",
        "第二步：演示一个与客户行业相近的客服场景。",
        "第三步：给出三档产品包，让客户选“先解决哪个场景”。",
        "第四步：承诺清晰交付物，而不是承诺万能 AI。",
        "第五步：用7天复盘推动续费、增购或转介绍。",
    ])

    add_heading(doc, "8. 需要继续补充的真实微信资料", 1)
    add_bullets(doc, [
        "客户备注名、标签、行业、已购产品或咨询产品。",
        "客户最原始的需求表达，最好保留原话。",
        "成交客户的付款/交付/复购信息。",
        "未成交客户停在哪一步：价格、信任、时间、决策人、产品不匹配。",
        "你自己的销售动作：发了什么案例、报价、演示、售后承诺，对方如何回应。",
    ])

    doc.save(OUT)


if __name__ == "__main__":
    main()
